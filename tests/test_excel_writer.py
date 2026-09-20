"""Excel writer tests."""

from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from rbws_extractor.excel.writer import RbwsExcelWriter, HEADERS
from rbws_extractor.models.findings import (
    AssetClass,
    ExtractionResult,
    Finding,
    ReportMetadata,
)


def test_excel_writer_headers_and_blank_judgment(tmp_path: Path):
    img_path = tmp_path / "photo.png"
    Image.new("RGB", (120, 80), color=(200, 40, 40)).save(img_path)

    result = ExtractionResult(
        metadata=ReportMetadata(facility="UBIT-GP", inspection_date=date(2024, 10, 9)),
        findings=[
            Finding(
                sn=1,
                asset_class=AssetClass.FIRE_WATER_PIPING,
                insp_date=date(2024, 10, 9),
                line_number='WF-010-8"-40-FAA-4',
                line_description="GP fire water header",
                anomaly_description="Heavy passive scale buildup (Photo #32-33)",
                pid_no="NGGP-AH-PDPID-0015210",
                photos_image_path=img_path,
            ).apply_judgment_blanks()
        ],
    )
    out = tmp_path / "RBWS.xlsx"
    RbwsExcelWriter().write(result, out, embed_images=True)
    assert out.exists()

    wb = load_workbook(out)
    ws = wb.active
    assert ws.title == "UBIT-GP"
    headers = [ws.cell(2, c).value for c in range(2, 18)]
    assert headers == HEADERS

    assert ws.cell(3, 2).value == 1
    assert ws.cell(3, 4).value == 'WF-010-8"-40-FAA-4'
    assert ws.cell(3, 9).value is None  # POF blank
    assert ws.cell(3, 16).value is None  # recommendations blank
    assert len(ws._images) == 1